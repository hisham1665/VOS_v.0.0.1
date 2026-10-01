"""Bounded extraction for non-image study documents.

The document loader deliberately keeps extraction local and deterministic.  It
never asks a model to interpret a file, and every limit that can discard data
is surfaced as a warning on the returned :class:`StudyDocument` (or as a
file-specific error in the study loader).
"""

from __future__ import annotations

import hashlib
import json
import re
import threading
import zipfile
from pathlib import Path
from typing import Iterable

from config import RUNTIME, SUPPORTED_DOCUMENT_SUFFIXES
from models import DocumentPage, StudyDocument

try:  # Optional at import time so image-only installations still work.
    import pypdfium2 as pdfium
except ImportError:  # pragma: no cover - exercised in minimal installations
    pdfium = None  # type: ignore[assignment]

try:  # pragma: no cover - import guard is environment dependent
    from docx import Document as _DocxDocument
except ImportError:
    _DocxDocument = None  # type: ignore[assignment,misc]

try:  # pragma: no cover - import guard is environment dependent
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore[assignment]


# PDFium uses a process-global native library.  In particular, opening a
# second document while another thread is rendering can crash on some builds.
# Keep the lock at module scope so all calls made by this module are guarded.
PDFIUM_LOCK = threading.RLock()

_TEXT_SUFFIXES = {
    ".txt", ".md", ".log", ".csv", ".tsv", ".json", ".ndjson", ".xml",
    ".html", ".htm", ".yaml", ".yml", ".rtf",
}
_BINARY_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


class DocumentLoadError(RuntimeError):
    """A file could not be safely opened or decoded."""


def _max_file_bytes() -> int:
    return max(1, int(RUNTIME.max_file_mb)) * 1024 * 1024


def _relative_path(path: Path, root: Path | None) -> str:
    """Return stable, portable provenance rather than an absolute local path."""
    path_abs = path.absolute()
    base = (root or path.parent).absolute()
    try:
        return path_abs.relative_to(base).as_posix()
    except ValueError:
        return path.name


def _document_id(path: Path, root: Path | None) -> str:
    # Relative paths are useful source references and are unique within a
    # study.  A direct load has the same intuitive id as its file name.
    return _relative_path(path, root)


def _new_document(path: Path, root: Path | None, content_hash: str) -> StudyDocument:
    suffix = path.suffix.lower()
    return StudyDocument(
        document_id=_document_id(path, root),
        source_path=_relative_path(path, root),
        content_hash=content_hash,
        format=suffix[1:] if suffix.startswith(".") else suffix,
    )


def _warn_once(values: list[str], warning: str) -> None:
    if warning and warning not in values:
        values.append(warning)


def _check_size(path: Path) -> int:
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise DocumentLoadError(f"cannot stat file: {exc}") from exc
    limit = _max_file_bytes()
    if size > limit:
        raise DocumentLoadError(
            f"file is {size / (1024 * 1024):.1f} MiB; maximum is "
            f"{RUNTIME.max_file_mb} MiB"
        )
    return size


def _read_bounded(path: Path) -> bytes:
    _check_size(path)
    try:
        return path.read_bytes()
    except OSError as exc:
        raise DocumentLoadError(f"cannot read file: {exc}") from exc


def _decode_text(raw: bytes) -> tuple[str, list[str]]:
    """Decode common clinical-export encodings and reject binary masquerading."""
    warnings: list[str] = []
    if raw.startswith(b"\xef\xbb\xbf"):
        try:
            return raw.decode("utf-8-sig"), warnings
        except UnicodeDecodeError as exc:  # pragma: no cover - impossible BOM case
            raise DocumentLoadError(f"invalid UTF-8 BOM text: {exc}") from exc

    if raw.startswith((b"\xff\xfe", b"\xfe\xff", b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        try:
            return raw.decode("utf-16"), warnings
        except UnicodeDecodeError as exc:
            raise DocumentLoadError(f"invalid UTF-16 text: {exc}") from exc

    try:
        return raw.decode("utf-8"), warnings
    except UnicodeDecodeError:
        pass

    # UTF-16 exports are occasionally emitted without a BOM.  A high NUL
    # ratio is a useful conservative signal; try both endian variants.
    if len(raw) >= 4 and len(raw) % 2 == 0:
        nul_ratio = raw.count(b"\x00") / len(raw)
        if nul_ratio >= 0.15:
            for encoding in ("utf-16-le", "utf-16-be"):
                try:
                    text = raw.decode(encoding)
                except UnicodeDecodeError:
                    continue
                if text and _binary_control_count(text) <= max(1, len(text) // 100):
                    warnings.append(f"decoded using {encoding} heuristic (no BOM)")
                    return text, warnings

    # Latin-1 is a total decoder, so validate its output before accepting it.
    # cp1252 gives the most useful result for old Windows exports.
    for encoding in ("cp1252", "latin-1"):
        text = raw.decode(encoding)
        if _binary_control_count(text) <= max(2, len(text) // 100):
            warnings.append(f"UTF-8 decoding failed; decoded using {encoding}")
            return text, warnings

    raise DocumentLoadError("file appears to be binary, not text")


def _binary_control_count(text: str) -> int:
    return len(_BINARY_CONTROL_RE.findall(text))


def _split_bounded_text(text: str, *, format_name: str, base_warnings: Iterable[str] = ()) -> tuple[list[DocumentPage], list[str]]:
    """Make bounded chunks without silently dropping over-budget text."""
    warnings = list(base_warnings)
    original_length = len(text)
    max_chars = max(1, int(RUNTIME.max_document_chars))
    chunk_chars = max(1, int(RUNTIME.document_chunk_chars))
    page_limit = max(1, int(RUNTIME.max_document_pages))

    if len(text) > max_chars:
        text = text[:max_chars]
        _warn_once(
            warnings,
            f"text truncated from {original_length} to {max_chars} characters "
            "by max_document_chars",
        )

    # Prefer row/line boundaries.  A single very long line is hard-split so a
    # malicious or malformed export cannot create an unbounded page.
    chunks: list[str] = []
    remaining = text
    while remaining:
        if len(remaining) <= chunk_chars:
            chunks.append(remaining)
            break
        cut = remaining.rfind("\n", 1, chunk_chars + 1)
        if cut <= 0:
            cut = chunk_chars
        chunks.append(remaining[:cut])
        remaining = remaining[cut:]
        if remaining.startswith("\n"):
            remaining = remaining[1:]

    if not chunks:
        chunks = [""]

    if len(chunks) > page_limit:
        represented = sum(len(part) for part in chunks[:page_limit])
        chunks = chunks[:page_limit]
        _warn_once(
            warnings,
            f"text truncated after {represented} characters because the document "
            f"exceeded max_document_pages={page_limit}",
        )

    pages = [
        DocumentPage(page_number=index, text=chunk, extraction=format_name)
        for index, chunk in enumerate(chunks, start=1)
    ]
    # Put truncation diagnostics on the affected final page as well as the
    # document, which makes them visible to consumers that inspect pages only.
    if warnings:
        pages[-1].warnings.extend(warnings)
    return pages, warnings


def _text_document(
    path: Path,
    root: Path | None,
    content_hash: str,
    text: str,
    format_name: str,
    warnings: Iterable[str] = (),
) -> StudyDocument:
    document = _new_document(path, root, content_hash)
    pages, page_warnings = _split_bounded_text(
        text, format_name=format_name, base_warnings=warnings
    )
    document.pages = pages
    document.page_count = len(pages)
    document.warnings.extend(page_warnings)
    return document


def _validate_office_zip(path: Path) -> list[str]:
    """Validate central-directory budgets before office libraries decompress."""
    _check_size(path)
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            max_bytes = _max_file_bytes()
            total_uncompressed = 0
            warnings: list[str] = []
            for info in infos:
                # A central directory can advertise a huge member without
                # allocating it.  Reject it before python-docx/openpyxl opens
                # any XML or media streams.
                if info.file_size > max_bytes:
                    raise DocumentLoadError(
                        "office archive contains an uncompressed member larger "
                        f"than the {RUNTIME.max_file_mb} MiB file budget"
                    )
                total_uncompressed += max(0, info.file_size)
                if info.compress_size and info.file_size / info.compress_size > 1000:
                    raise DocumentLoadError(
                        "office archive has an unsafe compression ratio (possible zip bomb)"
                    )
                if info.file_size and not info.compress_size:
                    raise DocumentLoadError(
                        "office archive contains an invalid zero-size compressed member"
                    )
            if total_uncompressed > max_bytes:
                raise DocumentLoadError(
                    "office archive uncompressed size exceeds the "
                    f"{RUNTIME.max_file_mb} MiB file budget"
                )
            names = [info.filename.replace("\\", "/") for info in infos]
            if any(name.startswith("xl/media/") for name in names):
                warnings.append("embedded spreadsheet images are not extracted")
            if any("embeddings/" in name for name in names):
                warnings.append("embedded office objects are not extracted")
            return warnings
    except DocumentLoadError:
        raise
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise DocumentLoadError(f"invalid office ZIP container: {exc}") from exc


def _load_pdf(path: Path, root: Path | None, content_hash: str) -> StudyDocument:
    document = _new_document(path, root, content_hash)
    if pdfium is None:
        document.warnings.append("PDF extraction unavailable: install pypdfium2")
        return document

    max_pages = max(1, int(RUNTIME.max_document_pages))
    image_edge = max(1, int(RUNTIME.document_image_edge))
    cache_dir = Path(RUNTIME.cache_dir) / "document_pages"
    cache_dir.mkdir(parents=True, exist_ok=True)
    try:
        # Do not move this lock into only the render call: opening, page
        # handles, text pages, and close all touch PDFium native state.
        with PDFIUM_LOCK:
            try:
                pdf = pdfium.PdfDocument(str(path))
            except Exception as exc:  # encrypted and malformed PDFs land here
                document.warnings.append(f"PDF could not be opened (encrypted or malformed): {exc}")
                return document
            try:
                try:
                    physical_count = len(pdf)
                except Exception as exc:
                    document.warnings.append(f"PDF page count failed: {exc}")
                    return document
                if physical_count == 0:
                    document.warnings.append("PDF contains no pages")
                    document.page_count = 0
                    return document
                if physical_count > max_pages:
                    document.warnings.append(
                        f"PDF has {physical_count} physical pages; extracted first "
                        f"{max_pages} due max_document_pages={max_pages}"
                    )

                text_chars = 0
                for physical_number in range(1, min(physical_count, max_pages) + 1):
                    page_warnings: list[str] = []
                    text = ""
                    image_path: Path | None = None
                    page = None
                    text_page = None
                    try:
                        page = pdf.get_page(physical_number - 1)
                        try:
                            text_page = page.get_textpage()
                            text = text_page.get_text_range() or ""
                        except Exception as exc:
                            page_warnings.append(f"native PDF text extraction failed: {exc}")

                        # Render every page, including pages with no native text.
                        try:
                            width, height = page.get_size()
                            longest = max(float(width), float(height), 1.0)
                            scale = image_edge / longest if longest > image_edge else 1.0
                            bitmap = page.render(scale=scale)
                            try:
                                rendered = bitmap.to_pil()
                                if rendered.mode not in ("RGB", "L"):
                                    rendered = rendered.convert("RGB")
                                image_path = cache_dir / (
                                    f"{content_hash[:24]}-page-{physical_number:04d}-edge-{image_edge}.jpg"
                                )
                                if not image_path.exists():
                                    temporary = image_path.with_suffix(".tmp.jpg")
                                    rendered.save(temporary, format="JPEG", quality=90, optimize=True)
                                    temporary.replace(image_path)
                            finally:
                                close = getattr(bitmap, "close", None)
                                if close:
                                    close()
                        except Exception as exc:
                            page_warnings.append(f"PDF page render failed: {exc}")
                    except Exception as exc:
                        page_warnings.append(f"PDF page {physical_number} failed: {exc}")
                    finally:
                        if text_page is not None:
                            close = getattr(text_page, "close", None)
                            if close:
                                close()
                        if page is not None:
                            close = getattr(page, "close", None)
                            if close:
                                close()

                    # Apply the document character budget to native text while
                    # retaining the physical page and its rendered image.
                    if text:
                        remaining = max(0, int(RUNTIME.max_document_chars) - text_chars)
                        if len(text) > remaining:
                            page_warnings.append(
                                "native PDF text truncated by max_document_chars"
                            )
                            text = text[:remaining]
                        text_chars += len(text)
                    if not text and not page_warnings:
                        page_warnings.append(
                            "no native PDF text; page rendered for visual extraction"
                        )
                    elif not text and image_path is not None and not any(
                        "native PDF text" in warning for warning in page_warnings
                    ):
                        page_warnings.append(
                            "no native PDF text; page rendered for visual extraction"
                        )

                    extraction = "pdf-native+rendered" if text else "pdf-rendered"
                    page = DocumentPage(
                        page_number=physical_number,
                        text=text,
                        image_path=str(image_path) if image_path else None,
                        extraction=extraction,
                        warnings=page_warnings,
                    )
                    document.pages.append(page)
                    for warning in page_warnings:
                        _warn_once(document.warnings, warning)

                if physical_count > max_pages:
                    # This is distinct from a text truncation warning: images
                    # for unprocessed physical pages cannot be represented in
                    # the bounded StudyDocument.
                    _warn_once(
                        document.warnings,
                        "remaining PDF pages were not loaded because of the page budget",
                    )
            finally:
                close = getattr(pdf, "close", None)
                if close:
                    close()
    except Exception as exc:  # defensive boundary around native PDFium calls
        _warn_once(document.warnings, f"PDF extraction failed: {exc}")

    # page_count is the physical count, while pages contains the bounded
    # successfully extracted subset.
    document.page_count = locals().get("physical_count", len(document.pages))
    return document


def _load_docx(path: Path, root: Path | None, content_hash: str) -> StudyDocument:
    warnings = _validate_office_zip(path)
    document = _new_document(path, root, content_hash)
    document.warnings.extend(warnings)
    if _DocxDocument is None:
        document.warnings.append("DOCX extraction unavailable: install python-docx")
        return document

    try:
        word_document = _DocxDocument(str(path))
        pieces: list[str] = []
        for paragraph in word_document.paragraphs:
            if paragraph.text:
                pieces.append(paragraph.text)
        for table in word_document.tables:
            for row in table.rows:
                pieces.append("\t".join(cell.text for cell in row.cells))
        text = "\n".join(pieces)
        # python-docx exposes inline images, while other relationships and
        # embedded objects are intentionally outside this text-only loader.
        if getattr(word_document, "inline_shapes", ()):
            _warn_once(document.warnings, "embedded DOCX images are not extracted")
        document.pages, page_warnings = _split_bounded_text(
            text, format_name="docx", base_warnings=document.warnings
        )
        document.warnings = page_warnings
        document.page_count = len(document.pages)
    except Exception as exc:
        document.warnings.append(f"DOCX extraction failed: {exc}")
    return document


def _load_xlsx(path: Path, root: Path | None, content_hash: str) -> StudyDocument:
    warnings = _validate_office_zip(path)
    document = _new_document(path, root, content_hash)
    document.warnings.extend(warnings)
    if openpyxl is None:
        document.warnings.append("XLSX extraction unavailable: install openpyxl")
        return document

    workbook = None
    values_workbook = None
    try:
        # Keep formulas rather than replacing them with None.  A second
        # read-only pass tells us whether a cached result is actually present.
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
        try:
            values_workbook = openpyxl.load_workbook(
                str(path), read_only=True, data_only=True
            )
        except Exception:
            values_workbook = None

        lines: list[str] = []
        for sheet_index, worksheet in enumerate(workbook.worksheets):
            lines.append(f"[Sheet: {worksheet.title}]")
            values_sheet = (
                values_workbook[worksheet.title] if values_workbook is not None else None
            )
            for row in worksheet.iter_rows():
                cells: list[str] = []
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    if cell.data_type == "f" or (
                        isinstance(value, str) and value.startswith("=")
                    ):
                        cached = None
                        if values_sheet is not None:
                            try:
                                cached = values_sheet[cell.coordinate].value
                            except Exception:
                                cached = None
                        if cached is None:
                            _warn_once(
                                document.warnings,
                                "spreadsheet formula has no cached result; formula text preserved",
                            )
                        else:
                            cells.append(f"{cell.coordinate}={value} [cached={cached}]")
                            continue
                    cells.append(f"{cell.coordinate}={value}")
                if cells:
                    lines.append("\t".join(cells))
        text = "\n".join(lines)
        document.pages, page_warnings = _split_bounded_text(
            text, format_name="xlsx", base_warnings=document.warnings
        )
        document.warnings = page_warnings
        document.page_count = len(document.pages)
    except Exception as exc:
        document.warnings.append(f"XLSX extraction failed: {exc}")
    finally:
        for loaded in (values_workbook, workbook):
            if loaded is not None:
                close = getattr(loaded, "close", None)
                if close:
                    close()
    return document


def load_document(path: str | Path, *, root: str | Path | None = None) -> StudyDocument:
    """Extract one supported document within configured safety budgets.

    The function returns a ``StudyDocument`` even for a format-specific
    failure, with a clear warning and zero pages.  ``load_study`` turns that
    warning into a file-specific skipped entry while retaining successful
    documents.
    """
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_DOCUMENT_SUFFIXES:
        raise DocumentLoadError(f"unsupported document extension '{suffix}'")
    root_path = Path(root) if root is not None else None
    try:
        raw = _read_bounded(source)
        content_hash = hashlib.sha256(raw).hexdigest()
    except DocumentLoadError as exc:
        document = _new_document(source, root_path, "")
        document.warnings.append(str(exc))
        return document

    try:
        if suffix == ".pdf":
            return _load_pdf(source, root_path, content_hash)
        if suffix in _TEXT_SUFFIXES:
            text, warnings = _decode_text(raw)
            if suffix == ".json":
                try:
                    parsed = json.loads(text)
                    text = json.dumps(parsed, ensure_ascii=False, indent=2)
                except (json.JSONDecodeError, TypeError, ValueError) as exc:
                    warnings = [*warnings, f"invalid JSON; preserved as text ({exc})"]
            # CSV/TSV are intentionally represented as text rows: this keeps
            # delimiters and values available to downstream consumers.
            return _text_document(source, root_path, content_hash, text, suffix[1:], warnings)
        if suffix == ".docx":
            return _load_docx(source, root_path, content_hash)
        if suffix == ".xlsx":
            return _load_xlsx(source, root_path, content_hash)
    except DocumentLoadError as exc:
        document = _new_document(source, root_path, content_hash)
        document.warnings.append(str(exc))
        return document
    except Exception as exc:  # per-file boundary; callers can continue a study
        document = _new_document(source, root_path, content_hash)
        document.warnings.append(f"{type(exc).__name__}: {exc}")
        return document

    # SUPPORTED_DOCUMENT_SUFFIXES is the source of truth, but keep a defensive
    # return for a future suffix added without an extractor.
    document = _new_document(source, root_path, content_hash)
    document.warnings.append(f"no extractor implemented for '{suffix}'")
    return document


# Friendly aliases for callers that use the verb from the module name.
extract_document = load_document


__all__ = [
    "DocumentLoadError",
    "PDFIUM_LOCK",
    "extract_document",
    "load_document",
]
